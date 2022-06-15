from typing import TYPE_CHECKING, Union
from .panel import GridPanel
import pandas as pd

from pptx.util import Inches, Pt

if TYPE_CHECKING:
    from grid_pptx.slide import GridSlide


class Table(GridPanel):
    def __init__(self, df: Union[pd.DataFrame, pd.Series], **kwargs):
        super().__init__(**kwargs)

        # convert df to dataframe if necessary
        if type(df) == pd.Series:
            self.df = df.to_frame()
        else:
            self.df = df

        self.minimize_height = True
        self.header = True
        self.first_col = False
        self.fontsize = 14

        # set any attributes that have been supplied in kwargs
        for k, v in kwargs.items():
            setattr(self, k, v)

        self.rows = len(df.index) + 1 if self.header else 0
        self.cols = len(df.columns)

        # if minimize height, then set a very small initial height so the table will be as compact as possible
        # note: must be done before creating table as cy is a required argument.
        if self.minimize_height:
            self.cy = Inches(0.5)

    def configure(self, gridslide: GridSlide) -> None:
        slide = gridslide.slide

        table = slide.shapes.add_table(
            self.rows, self.cols, self.x, self.y, self.cx, self.cy
        ).table

        # table formatting
        table.first_row = self.header
        table.first_col = self.first_col

        # populate header with df.columns if applicable
        row_count = 0
        if self.header:
            for i, val in enumerate(self.df.columns):
                table.cell(row_count, i).text = self.df.columns[i]
            row_count += 1

        # populate main table values
        for index, row in self.df.iterrows():
            for j, value in enumerate(row):
                table.cell(row_count, j).text = str(value)
                try:
                    # set font size
                    table.cell(row_count, j).text_frame.paragraphs[0].runs[0].font.size = Pt(self.fontsize)
                except IndexError:
                    # if passing "" as the value for the table, it will raise an error when trying to set
                    # characteristics of the font
                    pass
            row_count += 1
