from __future__ import annotations
from typing import TYPE_CHECKING, Union
from .panel import GridPanel
import pandas as pd

from pptx.util import Pt

# imports for type hints that would normally cause circular imports
if TYPE_CHECKING:
    from src.grid_pptx.slide import GridSlide


class Table(GridPanel):
    styles = {
        'no_style_no_grid': '{2D5ABB26-0587-4C30-8999-92F81FD0307C}',
        'themed_style_1_accent_1': '{3C2FFA5D-87B4-456A-9821-1D502468CF0F}',
        'themed_style_1_accent_2': '{284E427A-3D55-4303-BF80-6455036E1DE7}',
        'themed_style_1_accent_3': '{69C7853C-536D-4A76-A0AE-DD22124D55A5}',
        'themed_style_1_accent_4': '{775DCB02-9BB8-47FD-8907-85C794F793BA}',
        'themed_style_1_accent_5': '{35758FB7-9AC5-4552-8A53-C91805E547FA}',
        'themed_style_1_accent_6': '{08FB837D-C827-4EFA-A057-4D05807E0F7C}',
        'no_style_table_grid': '{5940675A-B579-460E-94D1-54222C63F5DA}',
        'themed_style_2_accent_1': '{D113A9D2-9D6B-4929-AA2D-F23B5EE8CBE7}',
        'themed_style_2_accent_2': '{18603FDC-E32A-4AB5-989C-0864C3EAD2B8}',
        'themed_style_2_accent_3': '{306799F8-075E-4A3A-A7F6-7FBC6576F1A4}',
        'themed_style_2_accent_4': '{E269D01E-BC32-4049-B463-5C60D7B0CCD2}',
        'themed_style_2_accent_5': '{327F97BB-C833-4FB7-BDE5-3F7075034690}',
        'themed_style_2_accent_6': '{638B1855-1B75-4FBE-930C-398BA8C253C6}',
        'light_style_1': '{9D7B26C5-4107-4FEC-AEDC-1716B250A1EF}',
        'light_style_1_accent_1': '{3B4B98B0-60AC-42C2-AFA5-B58CD77FA1E5}',
        'light_style_1_accent_2': '{0E3FDE45-AF77-4B5C-9715-49D594BDF05E}',
        'light_style_1_accent_3': '{C083E6E3-FA7D-4D7B-A595-EF9225AFEA82}',
        'light_style_1_accent_4': '{D27102A9-8310-4765-A935-A1911B00CA55}',
        'light_style_1_accent_5': '{5FD0F851-EC5A-4D38-B0AD-8093EC10F338}',
        'light_style_1_accent_6': '{68D230F3-CF80-4859-8CE7-A43EE81993B5}',
        'light_style_2': '{7E9639D4-E3E2-4D34-9284-5A2195B3D0D7}',
        'light_style_2_accent_1': '{69012ECD-51FC-41F1-AA8D-1B2483CD663E}',
        'light_style_2_accent_2': '{72833802-FEF1-4C79-8D5D-14CF1EAF98D9}',
        'light_style_2_accent_3': '{F2DE63D5-997A-4646-A377-4702673A728D}',
        'light_style_2_accent_4': '{17292A2E-F333-43FB-9621-5CBBE7FDCDCB}',
        'light_style_2_accent_5': '{5A111915-BE36-4E01-A7E5-04B1672EAD32}',
        'light_style_2_accent_6': '{912C8C85-51F0-491E-9774-3900AFEF0FD7}',
        'light_style_3': '{616DA210-FB5B-4158-B5E0-FEB733F419BA}',
        'light_style_3_accent_1': '{BC89EF96-8CEA-46FF-86C4-4CE0E7609802}',
        'light_style_3_accent_2': '{5DA37D80-6434-44D0-A028-1B22A696006F}',
        'light_style_3_accent_3': '{8799B23B-EC83-4686-B30A-512413B5E67A}',
        'light_style_3_accent_4': '{ED083AE6-46FA-4A59-8FB0-9F97EB10719F}',
        'light_style_3_accent_5': '{BDBED569-4797-4DF1-A0F4-6AAB3CD982D8}',
        'light_style_3_accent_6': '{E8B1032C-EA38-4F05-BA0D-38AFFFC7BED3}',
        'medium_style_1': '{793D81CF-94F2-401A-BA57-92F5A7B2D0C5}',
        'medium_style_1_accent_1': '{B301B821-A1FF-4177-AEE7-76D212191A09}',
        'medium_style_1_accent_2': '{9DCAF9ED-07DC-4A11-8D7F-57B35C25682E}',
        'medium_style_1_accent_3': '{1FECB4D8-DB02-4DC6-A0A2-4F2EBAE1DC90}',
        'medium_style_1_accent_4': '{1E171933-4619-4E11-9A3F-F7608DF75F80}',
        'medium_style_1_accent_5': '{FABFCF23-3B69-468F-B69F-88F6DE6A72F2}',
        'medium_style_1_accent_6': '{10A1B5D5-9B99-4C35-A422-299274C87663}',
        'medium_style_2': '{073A0DAA-6AF3-43AB-8588-CEC1D06C72B9}',
        'medium_style_2_accent_1': '{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}',
        'medium_style_2_accent_2': '{21E4AEA4-8DFA-4A89-87EB-49C32662AFE0}',
        'medium_style_2_accent_3': '{F5AB1C69-6EDB-4FF4-983F-18BD219EF322}',
        'medium_style_2_accent_4': '{00A15C55-8517-42AA-B614-E9B94910E393}',
        'medium_style_2_accent_5': '{7DF18680-E054-41AD-8BC1-D1AEF772440D}',
        'medium_style_2_accent_6': '{93296810-A885-4BE3-A3E7-6D5BEEA58F35}',
        'medium_style_3': '{8EC20E35-A176-4012-BC5E-935CFFF8708E}',
        'medium_style_3_accent_1': '{6E25E649-3F16-4E02-A733-19D2CDBF48F0}',
        'medium_style_3_accent_2': '{85BE263C-DBD7-4A20-BB59-AAB30ACAA65A}',
        'medium_style_3_accent_3': '{EB344D84-9AFB-497E-A393-DC336BA19D2E}',
        'medium_style_3_accent_4': '{EB9631B5-78F2-41C9-869B-9F39066F8104}',
        'medium_style_3_accent_5': '{74C1A8A3-306A-4EB7-A6B1-4F7E0EB9C5D6}',
        'medium_style_3_accent_6': '{2A488322-F2BA-4B5B-9748-0D474271808F}',
        'medium_style_4': '{D7AC3CCA-C797-4891-BE02-D94E43425B78}',
        'medium_style_4_accent_1': '{69CF1AB2-1976-4502-BF36-3FF5EA218861}',
        'medium_style_4_accent_2': '{8A107856-5554-42FB-B03E-39F5DBC370BA}',
        'medium_style_4_accent_3': '{0505E3EF-67EA-436B-97B2-0124C06EBD24}',
        'medium_style_4_accent_4': '{C4B1156A-380E-4F78-BDF5-A606A8083BF9}',
        'medium_style_4_accent_5': '{22838BEF-8BB2-4498-84A7-C5851F593DF1}',
        'medium_style_4_accent_6': '{16D9F66E-5EB9-4882-86FB-DCBF35E3C3E4}',
        'dark_style_1': '{E8034E78-7F5D-4C2E-B375-FC64B27BC917}',
        'dark_style_1_accent_1': '{125E5076-3810-47DD-B79F-674D7AD40C01}',
        'dark_style_1_accent_2': '{37CE84F3-28C3-443E-9E96-99CF82512B78}',
        'dark_style_1_accent_3': '{D03447BB-5D67-496B-8E87-E561075AD55C}',
        'dark_style_1_accent_4': '{E929F9F4-4A8F-4326-A1B4-22849713DDAB}',
        'dark_style_1_accent_5': '{8FD4443E-F989-4FC4-A0C8-D5A2AF1F390B}',
        'dark_style_1_accent_6': '{AF606853-7671-496A-8E4F-DF71F8EC918B}',
        'dark_style_2': '{5202B0CA-FC54-4496-8BCA-5EF66A818D29}',
        'dark_style_2_accent_1_2': '{0660B408-B3CF-4A94-85FC-2B1E0A45F4A2}',
        'dark_style_2_accent_3_4': '{91EBBBCC-DAD2-459C-BE2E-F6DE35CF9A28}',
        'dark_style_2_accent_5_6': '{46F890A9-2807-4EBB-B81D-B2AA78EC7F39}',
    }

    def __init__(self, df: Union[pd.DataFrame, pd.Series], style: str = 'medium_style_3_accent_1', **kwargs) -> None:
        """

        :param df:
        :param style:
        :param kwargs:
        """
        super().__init__(**kwargs)

        # convert df to dataframe if necessary
        if type(df) == pd.Series:
            self.df = df.to_frame()
        else:
            self.df = df

        self.minimize_height = True
        self.header = True
        self.first_col = False
        self.fontsize = 14
        self.style = style

        # set any attributes that have been supplied in kwargs
        for k, v in kwargs.items():
            setattr(self, k, v)

        self.rows = len(df.index) + 1 if self.header else 0
        self.cols = len(df.columns)

        # if minimize height, then set a very small initial height so the table will be as compact as possible
        # note: must be done before creating table as cy is a required argument.
        if self.minimize_height:
            self.height = 0.5

    def add_to_slide(self, gridslide: GridSlide) -> None:
        """

        :param gridslide:
        :return:
        """
        slide = gridslide.slide

        table_shape = slide.shapes.add_table(
            self.rows, self.cols, self.x, self.y, self.cx, self.cy
        )

        # table style -- ref https://github.com/scanny/python-pptx/issues/27
        tbl = table_shape._element.graphic.graphicData.tbl
        style_id = self.styles[self.style]
        tbl[0][-1].text = style_id

        table = table_shape.table

        # table formatting
        table.first_row = self.header
        table.first_col = self.first_col

        # populate header with df.columns if applicable
        row_count = 0
        if self.header:
            for i, val in enumerate(self.df.columns):
                table.cell(row_count, i).text = self.df.columns[i]
            row_count += 1

        # populate main table values
        for index, row in self.df.iterrows():
            for j, value in enumerate(row):
                table.cell(row_count, j).text = str(value)
                try:
                    # set font size
                    table.cell(row_count, j).text_frame.paragraphs[0].runs[0].font.size = Pt(self.fontsize)
                except IndexError:
                    # if passing "" as the value for the table, it will raise an error when trying to set
                    # characteristics of the font
                    pass
            row_count += 1
